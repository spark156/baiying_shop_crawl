import json
import shutil
import sys
import time
import csv
import html
import math
import asyncio
import urllib.request
import base64
import hashlib
import re
from io import BytesIO
from pathlib import Path

import history_store_json as history_store

from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

URL = "https://buyin.jinritemai.com/"

OUT = Path("outputs/baiying_json")
OUT.mkdir(parents=True, exist_ok=True)
PROFILE = Path("work/baiying_browser_profile")
DEBUG_LOG = Path("outputs/baiying_debug.log")
SCENE_DETAIL_JSON = Path("outputs/baiying_scene_detail.json")
SCENE_DETAIL_CSV = Path("outputs/baiying_scene_detail.csv")
SCENE_PRODUCTS_JSON = Path("outputs/baiying_scene_products.json")
SCENE_PRODUCTS_CSV = Path("outputs/baiying_scene_products.csv")
SCENE_PRODUCTS_HTML = Path("outputs/baiying_scene_products.html")
SCENE_PRODUCTS_XLSX = Path("outputs/baiying_scene_products.xlsx")
SCENE_PRODUCTS_XLSX_HD = Path("outputs/baiying_scene_products_hd.xlsx")
SCENE_PRODUCT_IMAGES = Path("outputs/baiying_product_images")
EXCEL_IMAGE_CACHE_SIZE = 768
EXCEL_IMAGE_DISPLAY_SIZE = 180
CDP_URL = "http://127.0.0.1:9222"


def get_arg_value(name, default=None):
    for index, arg in enumerate(sys.argv):
        if arg == name and index + 1 < len(sys.argv):
            return sys.argv[index + 1]
        if arg.startswith(name + "="):
            return arg.split("=", 1)[1]
    return default


def get_arg_int(name, default):
    value = get_arg_value(name)
    if value is None:
        return default
    try:
        return int(value)
    except ValueError:
        return default


def parse_shop_list(path):
    shops = []
    source = Path(path)
    for line_no, raw_line in enumerate(source.read_text(encoding="utf-8-sig").splitlines(), 1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue

        match = re.search(r"https?://\S+", line)
        if not match:
            append_debug(f"shop list line skipped without url: {line_no} {raw_line}")
            continue

        url = match.group(0).rstrip(",，")
        name = (line[: match.start()] + line[match.end() :]).strip(" \t,，")
        if not name:
            name = f"店铺{len(shops) + 1}"
        shops.append({"name": name, "url": url})

    if not shops:
        raise RuntimeError(f"店铺清单为空或没有可识别 URL：{source}")
    return shops


def append_debug(message):
    DEBUG_LOG.parent.mkdir(parents=True, exist_ok=True)
    with DEBUG_LOG.open("a", encoding="utf-8") as f:
        f.write(f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] {message}\n")


def reset_profile_if_requested():
    if "--reset-profile" not in sys.argv:
        return

    profile = PROFILE.resolve()
    work_root = Path("work").resolve()
    if work_root not in profile.parents:
        raise RuntimeError(f"refuse to delete profile outside work: {profile}")

    if PROFILE.exists():
        shutil.rmtree(str(PROFILE))
        print("reset profile:", PROFILE)


def find_system_browser():
    if "--chromium" in sys.argv:
        return None

    browser_names = ["edge"] if "--edge" in sys.argv else ["chrome", "edge"]
    candidates = {
        "chrome": [
            Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
            Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
            Path.home() / r"AppData\Local\Google\Chrome\Application\chrome.exe",
        ],
        "edge": [
            Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
            Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            Path.home() / r"AppData\Local\Microsoft\Edge\Application\msedge.exe",
        ],
    }

    for browser_name in browser_names:
        for executable in candidates[browser_name]:
            if executable.exists():
                return str(executable)
    return None


def clean_text(value):
    return "\n".join(line.strip() for line in value.splitlines() if line.strip())


def normalize_metric_value(value):
    lines = [line for line in clean_text(value).splitlines() if line != "查看流量结构"]
    return "\n".join(lines)


def parse_scene_detail_row(headers, cells, page_no):
    columns = {}
    for index, header in enumerate(headers):
        raw_value = cells[index] if index < len(cells) else ""
        columns[header] = normalize_metric_value(raw_value)

    live_info = columns.get("直播信息", "")
    live_lines = live_info.splitlines()
    row = {
        "页码": page_no,
        "直播时间": live_lines[0] if live_lines else "",
        "直播时长": live_lines[1].replace("时长：", "") if len(live_lines) > 1 else "",
    }
    row.update(columns)

    goods_lines = columns.get("带货商品", "").splitlines()
    if goods_lines:
        row["带货商品数"] = goods_lines[-1]
        if len(goods_lines) > 1 and goods_lines[0].startswith("+"):
            row["带货商品更多标记"] = goods_lines[0]

    return row


def extract_current_scene_table(page, page_no):
    table_data = page.evaluate(
        """
        () => {
            const tables = Array.from(document.querySelectorAll('table'));
            const table = tables.find((item) => {
                const text = item.innerText || '';
                return text.includes('直播信息') && text.includes('带货商品');
            });
            if (!table) {
                return null;
            }
            return {
                headers: Array.from(table.querySelectorAll('thead th')).map((th) => th.innerText.trim()),
                rows: Array.from(table.querySelectorAll('tbody tr')).map((tr) =>
                    Array.from(tr.querySelectorAll('td')).map((td) => td.innerText.trim())
                ),
            };
        }
        """
    )
    if not table_data:
        raise RuntimeError("没有找到场景分析的数据详情表格")

    return [
        parse_scene_detail_row(table_data["headers"], cells, page_no)
        for cells in table_data["rows"]
    ]


def get_scene_pagination_numbers(page):
    numbers = page.evaluate(
        """
        () => Array.from(document.querySelectorAll('li.auxo-pagination-item'))
            .filter((item) => !item.closest('.related-product-modal') && !item.closest('[role=dialog]'))
            .map((item) => Number((item.innerText || item.textContent || '').trim()))
            .filter((value) => Number.isInteger(value))
        """
    )
    return sorted(set(numbers)) or [1]


def get_active_scene_page(page):
    return page.evaluate(
        """
        () => {
            const active = Array.from(document.querySelectorAll('li.auxo-pagination-item-active'))
                .find((item) => !item.closest('.related-product-modal') && !item.closest('[role=dialog]'));
            if (!active) {
                return null;
            }
            const value = Number((active.innerText || active.textContent || '').trim());
            return Number.isInteger(value) ? value : null;
        }
        """
    )


def open_scene_analysis(page):
    page.evaluate(
        """
        () => {
            const tab = Array.from(document.querySelectorAll('.auxo-tabs-tab'))
                .find((item) => (item.innerText || '').trim() === '场景分析');
            if (tab) {
                tab.click();
            }
        }
        """
    )
    page.wait_for_selector("text=数据详情", timeout=10000)
    page.wait_for_selector("table", timeout=10000)
    page.wait_for_timeout(1000)


def switch_scene_page(page, page_no):
    current_page = get_active_scene_page(page)
    first_row_before = page.evaluate(
        """
        () => {
            const table = Array.from(document.querySelectorAll('table')).find((item) => {
                const text = item.innerText || '';
                return !item.closest('.related-product-modal')
                    && !item.closest('[role=dialog]')
                    && text.includes('直播信息')
                    && text.includes('带货商品');
            });
            return table?.querySelector('tbody tr')?.innerText || '';
        }
        """
    )

    page.evaluate(
        """
        ({ pageNo, currentPage }) => {
            const isMainPagination = (item) =>
                !item.closest('.related-product-modal') && !item.closest('[role=dialog]');
            let item = Array.from(document.querySelectorAll(`li.auxo-pagination-item-${pageNo}`))
                .find(isMainPagination);
            if (!item && currentPage && pageNo === currentPage + 1) {
                item = Array.from(document.querySelectorAll('li.auxo-pagination-next')).find(isMainPagination);
            }
            if (!item) {
                throw new Error(`scene pagination item not found: ${pageNo}`);
            }
            item.click();
        }
        """,
        {"pageNo": page_no, "currentPage": current_page},
    )

    try:
        page.wait_for_function(
            """
            (pageNo) => Array.from(document.querySelectorAll('li.auxo-pagination-item-active'))
                .filter((node) => !node.closest('.related-product-modal') && !node.closest('[role=dialog]'))
                .some((node) => (node.innerText || node.textContent || '').trim() === String(pageNo))
            """,
            arg=page_no,
            timeout=5000,
        )
    except PlaywrightTimeoutError:
        append_debug(f"pagination active timeout: {page_no}")

    try:
        page.wait_for_function(
            """
            (previousFirstRow) => {
                const table = Array.from(document.querySelectorAll('table')).find((item) => {
                    const text = item.innerText || '';
                    return !item.closest('.related-product-modal')
                        && !item.closest('[role=dialog]')
                        && text.includes('直播信息')
                        && text.includes('带货商品');
                });
                const currentFirstRow = table?.querySelector('tbody tr')?.innerText || '';
                return currentFirstRow && currentFirstRow !== previousFirstRow;
            }
            """,
            arg=first_row_before,
            timeout=8000,
        )
    except PlaywrightTimeoutError:
        append_debug(f"pagination first row unchanged: {page_no}")

    page.wait_for_timeout(1200)


def save_scene_detail(rows):
    SCENE_DETAIL_JSON.parent.mkdir(parents=True, exist_ok=True)
    SCENE_DETAIL_JSON.write_text(
        json.dumps(rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    fieldnames = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)

    with SCENE_DETAIL_CSV.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def format_price_cents(value):
    if value in (None, ""):
        return ""
    try:
        return "￥{:,.2f}".format(float(value) / 100)
    except (TypeError, ValueError):
        return str(value)


def format_yuan_amount(value):
    try:
        amount = float(value)
    except (TypeError, ValueError):
        return ""

    if amount == 0:
        return "0"
    if amount >= 10000:
        wan = amount / 10000
        return "{:g}万".format(wan)
    return "{:,.0f}".format(amount)


def format_yuan_range(low, high):
    try:
        low_value = float(low or 0)
        high_value = float(high or 0)
    except (TypeError, ValueError):
        return ""

    if low_value <= 0 and high_value <= 0:
        return "-"
    if low_value > 0 and high_value > 0:
        return "¥{}-{}".format(format_yuan_amount(low_value), format_yuan_amount(high_value))
    if low_value > 0:
        return "¥{}+".format(format_yuan_amount(low_value))
    return "¥{}".format(format_yuan_amount(high_value))


def has_settlement(product):
    return (
        product.get("sale_status_settle") == 0
        and (product.get("sale_low_settle") or product.get("sale_high_settle"))
    )


def close_related_product_modal(page):
    try:
        page.locator(".related-product-modal button").filter(has_text="知道了").click(timeout=2000)
    except Exception:
        page.evaluate(
            """
            () => {
                const modal = document.querySelector('.related-product-modal');
                if (!modal) {
                    return;
                }
                const okButton = Array.from(modal.querySelectorAll('button'))
                    .find((button) => (button.innerText || '').trim() === '知道了');
                if (okButton) {
                    okButton.click();
                }
            }
            """
        )

    try:
        page.wait_for_function(
            """
            () => {
                const modal = document.querySelector('.related-product-modal');
                if (!modal) {
                    return true;
                }
                const style = getComputedStyle(modal);
                const rect = modal.getBoundingClientRect();
                return style.display === 'none'
                    || style.visibility === 'hidden'
                    || rect.width === 0
                    || rect.height === 0;
            }
            """,
            timeout=3000,
        )
    except PlaywrightTimeoutError:
        append_debug("related product modal close timeout")


def extract_main_live_rows(page, page_no):
    rows = page.evaluate(
        """
        (pageNo) => {
            const table = Array.from(document.querySelectorAll('table')).find((item) => {
                const text = item.innerText || '';
                return !item.closest('.related-product-modal')
                    && text.includes('直播信息')
                    && text.includes('带货商品');
            });
            if (!table) {
                return [];
            }
            return Array.from(table.querySelectorAll('tbody tr')).map((tr, index) => {
                const cells = Array.from(tr.querySelectorAll('td')).map((td) => td.innerText.trim());
                const liveLines = (cells[0] || '').split('\\n').map((line) => line.trim()).filter(Boolean);
                const goodsLines = (cells[5] || '').split('\\n').map((line) => line.trim()).filter(Boolean);
                const productCount = Number(goodsLines[goodsLines.length - 1] || 0);
                return {
                    "直播页码": pageNo,
                    "直播行号": index + 1,
                    "直播ID": tr.getAttribute('data-row-key') || '',
                    "直播时间": liveLines[0] || '',
                    "直播时长": liveLines.length > 1 ? liveLines[1].replace('时长：', '') : '',
                    "观看人数": (cells[1] || '').replace('查看流量结构', '').trim(),
                    "人数峰值": cells[2] || '',
                    "主推类目": cells[3] || '',
                    "主推价格区间": cells[4] || '',
                    "带货商品数": Number.isFinite(productCount) ? productCount : 0,
                };
            });
        }
        """,
        page_no,
    )
    return rows


def click_main_live_product_button(page, row_index):
    page.evaluate(
        """
        (rowIndex) => {
            const table = Array.from(document.querySelectorAll('table')).find((item) => {
                const text = item.innerText || '';
                return !item.closest('.related-product-modal')
                    && text.includes('直播信息')
                    && text.includes('带货商品');
            });
            if (!table) {
                throw new Error('main live table not found');
            }
            const row = table.querySelectorAll('tbody tr')[rowIndex];
            if (!row) {
                throw new Error(`main live row not found: ${rowIndex}`);
            }
            const buttons = Array.from(row.querySelectorAll('td:last-child button'));
            const button = buttons.find((item) => /^\\d+$/.test((item.innerText || '').trim())) || buttons[buttons.length - 1];
            if (!button) {
                throw new Error(`product button not found: ${rowIndex}`);
            }
            button.click();
        }
        """,
        row_index,
    )


def wait_related_product_response(page, action, product_page):
    with page.expect_response(
        lambda resp: "scene_analysis/related_product_list" in resp.url
        and "page={}".format(product_page) in resp.url,
        timeout=15000,
    ) as response_info:
        action()
    return response_info.value.json()


def normalize_product_record(live_row, product, product_page, product_index):
    return {
        "店铺名称": live_row.get("店铺名称", ""),
        "直播页码": live_row.get("直播页码", ""),
        "直播行号": live_row.get("直播行号", ""),
        "直播ID": live_row.get("直播ID", ""),
        "直播时间": live_row.get("直播时间", ""),
        "直播时长": live_row.get("直播时长", ""),
        "商品分页": product_page,
        "商品序号": product_index,
        "商品ID": product.get("product_id", ""),
        "商品名称": product.get("title", ""),
        "商品链接": product.get("detail_url", ""),
        "商品图片": product.get("avatar", ""),
        "到手价": format_price_cents(product.get("real_price") or product.get("price")),
        "单场直播结算额": format_yuan_range(
            product.get("sale_low_settle"),
            product.get("sale_high_settle"),
        ),
        "到手价_raw": product.get("real_price") or product.get("price") or "",
        "单场直播结算额_low_raw": product.get("sale_low_settle") or 0,
        "单场直播结算额_high_raw": product.get("sale_high_settle") or 0,
    }


def product_detail_url(product_id):
    if not product_id:
        return ""
    return (
        "https://haohuo.jinritemai.com/ecommerce/trade/detail/index.html"
        "?id={}&origin_type=buyin_author_profile"
    ).format(product_id)


def extract_modal_product_rows(page, live_row, product_page):
    products = page.evaluate(
        """
        (productPage) => {
            const modal = document.querySelector('.related-product-modal');
            if (!modal) {
                return [];
            }
            const table = Array.from(modal.querySelectorAll('table')).find((item) => {
                const text = item.innerText || '';
                return text.includes('商品名称')
                    && text.includes('到手价')
                    && text.includes('单场直播结算额');
            });
            if (!table) {
                return [];
            }
            const normalizeUrl = (url) => {
                if (!url) {
                    return '';
                }
                const trimmed = url.split(',')[0].trim().split(' ')[0];
                if (trimmed.startsWith('//')) {
                    return `https:${trimmed}`;
                }
                return trimmed;
            };
            const getImageUrl = (tr) => {
                const image = tr.querySelector('img');
                const source = tr.querySelector('source');
                const candidates = [
                    image?.currentSrc,
                    image?.src,
                    image?.getAttribute('src'),
                    image?.getAttribute('srcset'),
                    source?.getAttribute('src'),
                    source?.getAttribute('srcset'),
                ];
                const realUrl = candidates.find((url) => url && !url.startsWith('data:'));
                return normalizeUrl(realUrl || '');
            };
            return Array.from(table.querySelectorAll('tbody tr')).map((tr, index) => {
                const cells = Array.from(tr.querySelectorAll('td')).map((td) => td.innerText.trim());
                const titleNode = tr.querySelector('.thumbnail-item-title');
                return {
                    product_page: productPage,
                    product_index: index + 1,
                    product_id: tr.getAttribute('data-row-key') || '',
                    title: titleNode?.getAttribute('title') || titleNode?.innerText?.trim() || cells[0] || '',
                    avatar: getImageUrl(tr),
                    price_text: cells[1] || '',
                    settlement_text: cells[2] || '',
                };
            });
        }
        """,
        product_page,
    )

    rows = []
    for product in products:
        settlement_text = clean_text(product.get("settlement_text", ""))
        if not settlement_text or settlement_text == "-":
            continue

        product_id = product.get("product_id", "")
        rows.append(
            {
                "店铺名称": live_row.get("店铺名称", ""),
                "直播页码": live_row.get("直播页码", ""),
                "直播行号": live_row.get("直播行号", ""),
                "直播ID": live_row.get("直播ID", ""),
                "直播时间": live_row.get("直播时间", ""),
                "直播时长": live_row.get("直播时长", ""),
                "商品分页": product.get("product_page", product_page),
                "商品序号": product.get("product_index", ""),
                "商品ID": product_id,
                "商品名称": product.get("title", ""),
                "商品链接": product_detail_url(product_id),
                "商品图片": product.get("avatar", ""),
                "到手价": clean_text(product.get("price_text", "")),
                "单场直播结算额": settlement_text,
                "到手价_raw": "",
                "单场直播结算额_low_raw": "",
                "单场直播结算额_high_raw": "",
            }
        )
    return rows


def collect_products_from_response(live_row, payload, product_page):
    products = payload.get("data") if isinstance(payload, dict) else []
    if not isinstance(products, list):
        return []

    rows = []
    for index, product in enumerate(products, 1):
        if has_settlement(product):
            rows.append(normalize_product_record(live_row, product, product_page, index))
    return rows


def capture_live_products(page, live_row, row_index):
    product_count = int(live_row.get("带货商品数") or 0)
    if product_count <= 0:
        return []

    product_page_count = max(1, int(math.ceil(product_count / 5)))
    click_main_live_product_button(page, row_index)
    page.wait_for_selector(".related-product-modal", timeout=10000)
    page.wait_for_selector(".related-product-modal table", timeout=10000)
    page.wait_for_timeout(800)

    rows = extract_modal_product_rows(page, live_row, 1)
    for product_page in range(2, product_page_count + 1):
        try:
            page.evaluate(
                """
                () => {
                    const modal = document.querySelector('.related-product-modal');
                    const next = modal?.querySelector('li.auxo-pagination-next');
                    if (!next || next.getAttribute('aria-disabled') === 'true') {
                        throw new Error('product next page is not available');
                    }
                    next.click();
                }
                """
            )
            page.wait_for_timeout(900)
        except Exception as exc:
            append_debug(
                "related product page failed: live={} product_page={} error={}".format(
                    live_row.get("直播时间", ""),
                    product_page,
                    exc,
                )
            )
            break
        rows.extend(extract_modal_product_rows(page, live_row, product_page))

    close_related_product_modal(page)
    return rows


def download_product_image(image_url, product_id):
    if not image_url:
        return None

    SCENE_PRODUCT_IMAGES.mkdir(parents=True, exist_ok=True)
    key = product_id or image_url
    digest = hashlib.md5(key.encode("utf-8")).hexdigest()
    image_path = SCENE_PRODUCT_IMAGES / f"{digest}_{EXCEL_IMAGE_CACHE_SIZE}.png"
    if image_path.exists():
        return image_path

    cached_versions = sorted(
        SCENE_PRODUCT_IMAGES.glob(f"{digest}_*.png"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if cached_versions:
        return cached_versions[0]

    request = urllib.request.Request(
        image_url,
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0 Safari/537.36"
            ),
            "Referer": "https://buyin.jinritemai.com/",
        },
    )
    # System proxy settings are often intended only for browsers. A stale local
    # proxy makes every CDN image fail with WinError 10061, so fetch images directly.
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    with opener.open(request, timeout=20) as response:
        image_bytes = response.read()

    from PIL import Image as PilImage

    with PilImage.open(BytesIO(image_bytes)) as image:
        resampling = getattr(getattr(PilImage, "Resampling", PilImage), "LANCZOS", 1)
        image.thumbnail((EXCEL_IMAGE_CACHE_SIZE, EXCEL_IMAGE_CACHE_SIZE), resampling)
        if image.mode not in ("RGB", "RGBA"):
            image = image.convert("RGBA")
        image.save(image_path, "PNG")

    return image_path


def save_scene_products_xlsx(rows, output_path=None):
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as ExcelImage
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        append_debug(f"xlsx dependency unavailable: {type(exc).__name__}: {exc}")
        return None

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品明细"

    headers = [
        "店铺名称",
        "直播页码",
        "直播行号",
        "直播时间",
        "直播时长",
        "商品图片",
        "商品名称",
        "商品链接",
        "商品图片链接",
        "到手价",
        "单场直播结算额",
        "商品ID",
        "直播ID",
    ]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="DCEBFF")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    column_widths = {
        "A": 24,
        "B": 10,
        "C": 10,
        "D": 22,
        "E": 18,
        "F": 28,
        "G": 36,
        "H": 48,
        "I": 48,
        "J": 16,
        "K": 18,
        "L": 24,
        "M": 24,
    }
    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width

    image_column = headers.index("商品图片") + 1
    image_paths = []
    image_url_count = 0
    embedded_image_count = 0
    for row_index, row in enumerate(rows, start=2):
        values = [
            row.get("店铺名称", ""),
            row.get("直播页码", ""),
            row.get("直播行号", ""),
            row.get("直播时间", ""),
            row.get("直播时长", ""),
            "",
            row.get("商品名称", ""),
            row.get("商品链接", ""),
            row.get("商品图片", ""),
            row.get("到手价", ""),
            row.get("单场直播结算额", ""),
            row.get("商品ID", ""),
            row.get("直播ID", ""),
        ]
        sheet.append(values)
        sheet.row_dimensions[row_index].height = 138

        for cell in sheet[row_index]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        product_link = row.get("商品链接", "")
        if product_link:
            link_cell = sheet.cell(row=row_index, column=headers.index("商品链接") + 1)
            link_cell.hyperlink = product_link
            link_cell.style = "Hyperlink"

        image_url = row.get("商品图片", "")
        if image_url:
            image_url_count += 1
            image_link_cell = sheet.cell(row=row_index, column=headers.index("商品图片链接") + 1)
            image_link_cell.hyperlink = image_url
            image_link_cell.style = "Hyperlink"

        try:
            image_path = download_product_image(image_url, row.get("商品ID", ""))
        except Exception as exc:
            append_debug(f"download product image failed: {image_url} {exc}")
            image_path = None

        if image_path:
            image_paths.append(image_path)
            excel_image = ExcelImage(str(image_path))
            excel_image.width = EXCEL_IMAGE_DISPLAY_SIZE
            excel_image.height = EXCEL_IMAGE_DISPLAY_SIZE
            anchor = f"{get_column_letter(image_column)}{row_index}"
            sheet.add_image(excel_image, anchor)
            embedded_image_count += 1

    if image_url_count and embedded_image_count == 0:
        raise RuntimeError(
            "商品图片下载全部失败，Excel 未生成。请确认电脑可以直接访问商品图片链接后重试。"
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    target_path = Path(output_path) if output_path else SCENE_PRODUCTS_XLSX
    fallback_path = (
        target_path.with_name(target_path.stem + "_hd" + target_path.suffix)
        if output_path
        else SCENE_PRODUCTS_XLSX_HD
    )
    target_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(target_path)
        return target_path
    except PermissionError:
        workbook.save(fallback_path)
        return fallback_path


def save_scene_products(rows):
    SCENE_PRODUCTS_JSON.parent.mkdir(parents=True, exist_ok=True)
    SCENE_PRODUCTS_JSON.write_text(
        json.dumps(rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    fieldnames = [
        "店铺名称",
        "店铺网址",
        "直播日期",
        "直播页码",
        "直播行号",
        "直播时间",
        "直播时长",
        "直播ID",
        "商品分页",
        "商品序号",
        "商品ID",
        "商品名称",
        "商品链接",
        "商品图片",
        "到手价",
        "单场直播结算额",
        "到手价_raw",
        "单场直播结算额_low_raw",
        "单场直播结算额_high_raw",
    ]
    with SCENE_PRODUCTS_CSV.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    html_rows = []
    for row in rows:
        product_name = html.escape(str(row.get("商品名称", "")))
        product_url = html.escape(str(row.get("商品链接", "")))
        image_url = html.escape(str(row.get("商品图片", "")))
        html_rows.append(
            """
            <tr>
              <td>{shop_name}</td>
              <td>{live_time}</td>
              <td><a href="{product_url}" target="_blank" rel="noreferrer">{product_name}</a></td>
              <td><img src="{image_url}" loading="lazy" /></td>
              <td>{price}</td>
              <td>{settlement}</td>
            </tr>
            """.format(
                shop_name=html.escape(str(row.get("店铺名称", ""))),
                live_time=html.escape(str(row.get("直播时间", ""))),
                product_url=product_url,
                product_name=product_name,
                image_url=image_url,
                price=html.escape(str(row.get("到手价", ""))),
                settlement=html.escape(str(row.get("单场直播结算额", ""))),
            )
        )

    SCENE_PRODUCTS_HTML.write_text(
        """<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <title>场景分析商品明细</title>
  <style>
    body {{ font-family: Arial, "Microsoft YaHei", sans-serif; margin: 24px; color: #1f2937; }}
    table {{ border-collapse: collapse; width: 100%; }}
    th, td {{ border: 1px solid #d9dee8; padding: 8px 10px; text-align: left; vertical-align: middle; }}
    th {{ background: #f3f6fb; position: sticky; top: 0; }}
    img {{ width: 72px; height: 72px; object-fit: cover; border-radius: 4px; }}
    a {{ color: #1459ff; text-decoration: none; }}
  </style>
</head>
<body>
  <h1>场景分析商品明细</h1>
  <table>
    <thead>
      <tr>
        <th>店铺名称</th>
        <th>直播时间</th>
        <th>商品名称/链接</th>
        <th>图片</th>
        <th>到手价</th>
        <th>单场直播结算额</th>
      </tr>
    </thead>
    <tbody>
      {rows}
    </tbody>
  </table>
</body>
</html>
""".format(rows="\n".join(html_rows)),
        encoding="utf-8",
    )
    save_scene_products_xlsx(rows)


def capture_scene_products(page):
    close_related_product_modal(page)
    open_scene_analysis(page)
    close_related_product_modal(page)

    max_live_pages = get_arg_int("--max-pages", 2)
    max_live_rows = get_arg_int("--max-live-rows", 0)
    available_pages = get_scene_pagination_numbers(page)
    target_pages = [page_no for page_no in available_pages if page_no <= max_live_pages]
    rows = []

    for page_no in target_pages:
        switch_scene_page(page, page_no)
        live_rows = extract_main_live_rows(page, page_no)
        if max_live_rows > 0:
            live_rows = live_rows[:max_live_rows]
        print(f"live page {page_no}: {len(live_rows)} live rows", flush=True)

        for row_index, live_row in enumerate(live_rows):
            try:
                product_rows = capture_live_products(page, live_row, row_index)
            except Exception as exc:
                append_debug(
                    "capture live products failed: page={} row={} live={} error={}".format(
                        page_no,
                        row_index + 1,
                        live_row.get("直播时间", ""),
                        exc,
                    )
                )
                close_related_product_modal(page)
                product_rows = []

            print(
                "  live row {} {}: {} settled products".format(
                    row_index + 1,
                    live_row.get("直播时间", ""),
                    len(product_rows),
                ),
                flush=True,
            )
            rows.extend(product_rows)

    save_scene_products(rows)
    print("scene product rows:", len(rows), flush=True)
    print("scene product json:", SCENE_PRODUCTS_JSON, flush=True)
    print("scene product csv:", SCENE_PRODUCTS_CSV, flush=True)
    print("scene product html:", SCENE_PRODUCTS_HTML, flush=True)
    print("scene product xlsx:", SCENE_PRODUCTS_XLSX, flush=True)
    return rows


class RawCdpPage:
    def __init__(self, ws_url):
        self.ws_url = ws_url
        self.conn = None
        self.next_id = 1
        self.events = []
        self.response_urls = {}
        self.finished_requests = set()

    async def __aenter__(self):
        import websockets

        self.conn = await websockets.connect(self.ws_url, max_size=20 * 1024 * 1024)
        await self.send("Page.bringToFront")
        await self.send("Network.enable")
        return self

    async def __aexit__(self, exc_type, exc, tb):
        if self.conn:
            await self.conn.close()

    async def send(self, method, params=None):
        message_id = self.next_id
        self.next_id += 1
        await self.conn.send(
            json.dumps(
                {
                    "id": message_id,
                    "method": method,
                    "params": params or {},
                }
            )
        )
        while True:
            message = json.loads(await self.conn.recv())
            if message.get("id") != message_id:
                self.record_event(message)
                continue
            if "error" in message:
                raise RuntimeError(message["error"])
            return message.get("result", {})

    def record_event(self, message):
        self.events.append(message)
        method = message.get("method")
        params = message.get("params", {})
        if method == "Network.responseReceived":
            response = params.get("response", {})
            request_id = params.get("requestId")
            if request_id:
                self.response_urls[request_id] = response.get("url", "")
        elif method == "Network.loadingFinished":
            request_id = params.get("requestId")
            if request_id:
                self.finished_requests.add(request_id)

    async def pump_events(self, seconds):
        deadline = time.time() + seconds
        while time.time() < deadline:
            try:
                message = await asyncio.wait_for(self.conn.recv(), timeout=max(0.05, deadline - time.time()))
            except asyncio.TimeoutError:
                break
            self.record_event(json.loads(message))

    async def get_json_response_body(self, request_id):
        body_result = await self.send("Network.getResponseBody", {"requestId": request_id})
        body = body_result.get("body", "")
        if body_result.get("base64Encoded"):
            body = base64.b64decode(body).decode("utf-8", errors="replace")
        return json.loads(body)

    async def get_related_product_payloads_since(self, start_index, product_page=None):
        payloads = []
        for message in self.events[start_index:]:
            if message.get("method") != "Network.responseReceived":
                continue
            params = message.get("params", {})
            request_id = params.get("requestId")
            url = params.get("response", {}).get("url", "")
            if "scene_analysis/related_product_list" not in url:
                continue
            if product_page is not None and "page={}".format(product_page) not in url:
                continue
            try:
                payloads.append(await self.get_json_response_body(request_id))
            except Exception as exc:
                append_debug(f"raw get response body failed: {url} {exc}")
        return payloads

    async def evaluate(self, expression):
        result = await self.send(
            "Runtime.evaluate",
            {
                "expression": expression,
                "returnByValue": True,
                "awaitPromise": True,
            },
        )
        if "exceptionDetails" in result:
            raise RuntimeError(result["exceptionDetails"])
        return result.get("result", {}).get("value")


def get_raw_cdp_page_ws(cdp_url, allow_any_page=False):
    with urllib.request.urlopen(cdp_url.rstrip("/") + "/json/list", timeout=5) as response:
        targets = json.loads(response.read().decode("utf-8"))
    for target in targets:
        if target.get("type") == "page" and "buyin.jinritemai.com" in target.get("url", ""):
            return target["webSocketDebuggerUrl"]
    if allow_any_page:
        for target in targets:
            if target.get("type") == "page":
                return target["webSocketDebuggerUrl"]
    raise RuntimeError("没有找到 buyin.jinritemai.com 页面，请确认调试 Chrome 里已经打开达人主页")


def js_call(function_body, arg=None):
    return "({})({})".format(function_body, json.dumps(arg, ensure_ascii=False))


async def raw_sleep(seconds):
    await asyncio.sleep(seconds)


async def raw_wait_for_scene_tab(page, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        has_tab = await page.evaluate(
            """
            (() => Array.from(document.querySelectorAll('.auxo-tabs-tab'))
                .some((item) => (item.innerText || '').trim() === '场景分析'))()
            """
        )
        if has_tab:
            return True
        await raw_sleep(0.5)
    return False


async def raw_navigate_to_url(page, url):
    await raw_close_related_product_modal(page)
    await page.send("Page.navigate", {"url": url})
    await raw_sleep(2.5)

    deadline = time.time() + 45
    while time.time() < deadline:
        state = await page.evaluate(
            """
            (() => ({
                href: location.href,
                ready: document.readyState,
                text: document.body ? document.body.innerText.slice(0, 200) : ''
            }))()
            """
        )
        if state and state.get("ready") in ("interactive", "complete"):
            if "buyin.jinritemai.com" in state.get("href", ""):
                break
        await raw_sleep(0.5)

    if not await raw_wait_for_scene_tab(page, timeout=35):
        raise RuntimeError(f"页面没有加载出“场景分析”标签：{url}")


async def raw_close_related_product_modal(page):
    button_rect = await page.evaluate(
        """
        (() => {
            const modal = document.querySelector('.related-product-modal');
            if (!modal) {
                return null;
            }
            const okButton = Array.from(modal.querySelectorAll('button'))
                .find((button) => (button.innerText || '').trim() === '知道了');
            if (okButton) {
                const rect = okButton.getBoundingClientRect();
                return {
                    x: rect.left + rect.width / 2,
                    y: rect.top + rect.height / 2,
                };
            }
            return null;
        })()
        """
    )
    if button_rect:
        await page.send(
            "Input.dispatchMouseEvent",
            {
                "type": "mouseMoved",
                "x": button_rect["x"],
                "y": button_rect["y"],
            },
        )
        await page.send(
            "Input.dispatchMouseEvent",
            {
                "type": "mousePressed",
                "x": button_rect["x"],
                "y": button_rect["y"],
                "button": "left",
                "clickCount": 1,
            },
        )
        await page.send(
            "Input.dispatchMouseEvent",
            {
                "type": "mouseReleased",
                "x": button_rect["x"],
                "y": button_rect["y"],
                "button": "left",
                "clickCount": 1,
            },
        )
    await raw_sleep(0.6)
    await page.send(
        "Input.dispatchKeyEvent",
        {
            "type": "keyDown",
            "key": "Escape",
            "code": "Escape",
            "windowsVirtualKeyCode": 27,
        },
    )
    await page.send(
        "Input.dispatchKeyEvent",
        {
            "type": "keyUp",
            "key": "Escape",
            "code": "Escape",
            "windowsVirtualKeyCode": 27,
        },
    )
    await raw_sleep(0.3)
    still_visible = await page.evaluate(
        """
        (() => {
            const modal = document.querySelector('.related-product-modal');
            if (!modal) {
                return false;
            }
            const style = getComputedStyle(modal);
            const rect = modal.getBoundingClientRect();
            return style.display !== 'none'
                && style.visibility !== 'hidden'
                && rect.width > 0
                && rect.height > 0;
        })()
        """
    )
    if still_visible:
        await page.evaluate(
            """
            (() => {
                const roots = new Set();
                document.querySelectorAll('.related-product-modal').forEach((modal) => {
                    const root = modal.closest('.auxo-modal-root') || modal.closest('[role=dialog]') || modal;
                    roots.add(root);
                });
                roots.forEach((root) => root.remove());
                document.body.classList.remove('ant-scrolling-effect');
                document.body.style.removeProperty('overflow');
                document.body.style.removeProperty('width');
                return true;
            })()
            """
        )
        await raw_sleep(0.2)


async def raw_open_scene_analysis(page):
    await raw_close_related_product_modal(page)
    if not await raw_wait_for_scene_tab(page, timeout=35):
        raise RuntimeError("页面没有找到“场景分析”标签")
    await page.evaluate(
        """
        (() => {
            const tab = Array.from(document.querySelectorAll('.auxo-tabs-tab'))
                .find((item) => (item.innerText || '').trim() === '场景分析');
            if (tab) {
                tab.click();
            }
            return true;
        })()
        """
    )
    await raw_sleep(1.2)


async def raw_get_scene_pagination_numbers(page):
    numbers = await page.evaluate(
        """
        (() => Array.from(document.querySelectorAll('li.auxo-pagination-item'))
            .filter((item) => !item.closest('.related-product-modal') && !item.closest('[role=dialog]'))
            .map((item) => Number((item.innerText || item.textContent || '').trim()))
            .filter((value) => Number.isInteger(value)))()
        """
    )
    return sorted(set(numbers or [])) or [1]


async def raw_switch_scene_page(page, page_no):
    await raw_close_related_product_modal(page)
    await page.evaluate(
        js_call(
            """
            ({ pageNo }) => {
                const isMainPagination = (item) =>
                    !item.closest('.related-product-modal') && !item.closest('[role=dialog]');
                const item = Array.from(document.querySelectorAll(`li.auxo-pagination-item-${pageNo}`))
                    .find(isMainPagination);
                if (!item) {
                    throw new Error(`scene pagination item not found: ${pageNo}`);
                }
                item.click();
                return true;
            }
            """,
            {"pageNo": page_no},
        )
    )
    await raw_sleep(1.2)


async def raw_extract_main_live_rows(page, page_no):
    return await page.evaluate(
        js_call(
            """
            (pageNo) => {
                const table = Array.from(document.querySelectorAll('table')).find((item) => {
                    const text = item.innerText || '';
                    return !item.closest('.related-product-modal')
                        && !item.closest('[role=dialog]')
                        && text.includes('直播信息')
                        && text.includes('带货商品');
                });
                if (!table) {
                    return [];
                }
                return Array.from(table.querySelectorAll('tbody tr')).map((tr, index) => {
                    const cells = Array.from(tr.querySelectorAll('td')).map((td) => td.innerText.trim());
                    const liveLines = (cells[0] || '').split('\\n').map((line) => line.trim()).filter(Boolean);
                    const goodsLines = (cells[5] || '').split('\\n').map((line) => line.trim()).filter(Boolean);
                    const productCount = Number(goodsLines[goodsLines.length - 1] || 0);
                    return {
                        "直播页码": pageNo,
                        "直播行号": index + 1,
                        "直播ID": tr.getAttribute('data-row-key') || '',
                        "直播时间": liveLines[0] || '',
                        "直播时长": liveLines.length > 1 ? liveLines[1].replace('时长：', '') : '',
                        "观看人数": (cells[1] || '').replace('查看流量结构', '').trim(),
                        "人数峰值": cells[2] || '',
                        "主推类目": cells[3] || '',
                        "主推价格区间": cells[4] || '',
                        "带货商品数": Number.isFinite(productCount) ? productCount : 0,
                    };
                });
            }
            """,
            page_no,
        )
    )


async def raw_click_live_product_button(page, row_index):
    await page.evaluate(
        js_call(
            """
            (rowIndex) => {
                const table = Array.from(document.querySelectorAll('table')).find((item) => {
                    const text = item.innerText || '';
                    return !item.closest('.related-product-modal')
                        && !item.closest('[role=dialog]')
                        && text.includes('直播信息')
                        && text.includes('带货商品');
                });
                if (!table) {
                    throw new Error('main live table not found');
                }
                const row = table.querySelectorAll('tbody tr')[rowIndex];
                if (!row) {
                    throw new Error(`main live row not found: ${rowIndex}`);
                }
                const buttons = Array.from(row.querySelectorAll('td:last-child button'));
                const button = buttons.find((item) => /^\\d+$/.test((item.innerText || '').trim())) || buttons[buttons.length - 1];
                if (!button) {
                    throw new Error(`product button not found: ${rowIndex}`);
                }
                button.click();
                return true;
            }
            """,
            row_index,
        )
    )


async def raw_wait_for_product_modal(page, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        visible = await page.evaluate(
            """
            (() => {
                const modal = document.querySelector('.related-product-modal');
                if (!modal) {
                    return false;
                }
                const rect = modal.getBoundingClientRect();
                const style = getComputedStyle(modal);
                return style.display !== 'none'
                    && style.visibility !== 'hidden'
                    && rect.width > 0
                    && rect.height > 0
                    && !!modal.querySelector('table tbody tr');
            })()
            """
        )
        if visible:
            return
        await raw_sleep(0.3)
    raise TimeoutError("等待商品弹窗超时")


async def raw_extract_modal_product_rows(page, live_row, product_page):
    products = await page.evaluate(
        js_call(
            """
            (productPage) => {
                const modal = document.querySelector('.related-product-modal');
                if (!modal) {
                    return [];
                }
                const table = Array.from(modal.querySelectorAll('table')).find((item) => {
                    const text = item.innerText || '';
                    return text.includes('商品名称')
                        && text.includes('到手价')
                        && text.includes('单场直播结算额');
                });
                if (!table) {
                    return [];
                }
                const normalizeUrl = (url) => {
                    if (!url) {
                        return '';
                    }
                    const trimmed = url.split(',')[0].trim().split(' ')[0];
                    if (trimmed.startsWith('//')) {
                        return `https:${trimmed}`;
                    }
                    return trimmed;
                };
                const getImageUrl = (tr) => {
                    const image = tr.querySelector('img');
                    const source = tr.querySelector('source');
                    const candidates = [
                        image?.currentSrc,
                        image?.src,
                        image?.getAttribute('src'),
                        image?.getAttribute('srcset'),
                        source?.getAttribute('src'),
                        source?.getAttribute('srcset'),
                    ];
                    const realUrl = candidates.find((url) => url && !url.startsWith('data:'));
                    return normalizeUrl(realUrl || '');
                };
                return Array.from(table.querySelectorAll('tbody tr')).map((tr, index) => {
                    const cells = Array.from(tr.querySelectorAll('td')).map((td) => td.innerText.trim());
                    const titleNode = tr.querySelector('.thumbnail-item-title');
                    return {
                        product_page: productPage,
                        product_index: index + 1,
                        product_id: tr.getAttribute('data-row-key') || '',
                        title: titleNode?.getAttribute('title') || titleNode?.innerText?.trim() || cells[0] || '',
                        avatar: getImageUrl(tr),
                        price_text: cells[1] || '',
                        settlement_text: cells[2] || '',
                    };
                });
            }
            """,
            product_page,
        )
    )

    rows = []
    for product in products or []:
        settlement_text = clean_text(product.get("settlement_text", ""))
        if not settlement_text or settlement_text == "-":
            continue
        product_id = product.get("product_id", "")
        rows.append(
            {
                "店铺名称": live_row.get("店铺名称", ""),
                "直播页码": live_row.get("直播页码", ""),
                "直播行号": live_row.get("直播行号", ""),
                "直播ID": live_row.get("直播ID", ""),
                "直播时间": live_row.get("直播时间", ""),
                "直播时长": live_row.get("直播时长", ""),
                "商品分页": product.get("product_page", product_page),
                "商品序号": product.get("product_index", ""),
                "商品ID": product_id,
                "商品名称": product.get("title", ""),
                "商品链接": product_detail_url(product_id),
                "商品图片": product.get("avatar", ""),
                "到手价": clean_text(product.get("price_text", "")),
                "单场直播结算额": settlement_text,
                "到手价_raw": "",
                "单场直播结算额_low_raw": "",
                "单场直播结算额_high_raw": "",
            }
        )
    return rows


async def raw_click_product_next_page(page):
    return await page.evaluate(
        """
        (() => {
            const modal = document.querySelector('.related-product-modal');
            const next = modal?.querySelector('li.auxo-pagination-next');
            if (!next || next.getAttribute('aria-disabled') === 'true') {
                return false;
            }
            next.click();
            return true;
        })()
        """
    )


async def raw_capture_live_products(page, live_row, row_index):
    product_count = int(live_row.get("带货商品数") or 0)
    if product_count <= 0:
        return []

    product_page_count = max(1, int(math.ceil(product_count / 5)))
    event_start = len(page.events)
    await raw_click_live_product_button(page, row_index)
    await raw_wait_for_product_modal(page)
    await page.pump_events(1.0)

    payloads = await page.get_related_product_payloads_since(event_start, 1)
    if payloads:
        rows = collect_products_from_response(live_row, payloads[-1], 1)
    else:
        rows = await raw_extract_modal_product_rows(page, live_row, 1)
    for product_page in range(2, product_page_count + 1):
        event_start = len(page.events)
        has_next = await raw_click_product_next_page(page)
        if not has_next:
            break
        await page.pump_events(1.0)
        payloads = await page.get_related_product_payloads_since(event_start, product_page)
        if payloads:
            rows.extend(collect_products_from_response(live_row, payloads[-1], product_page))
        else:
            rows.extend(await raw_extract_modal_product_rows(page, live_row, product_page))

    await raw_close_related_product_modal(page)
    return rows


def emit_progress(progress, message, **details):
    if progress:
        progress({"message": message, **details})
    else:
        print(message, flush=True)


async def raw_capture_scene_products_from_page(
    page,
    shop_name="",
    shop_url="",
    start_date="",
    end_date="",
    max_live_pages=None,
    force=False,
    progress=None,
):
    await raw_open_scene_analysis(page)
    await raw_close_related_product_modal(page)

    if max_live_pages is None:
        max_live_pages = get_arg_int("--max-pages", 2)
    max_live_rows = get_arg_int("--max-live-rows", 0)
    available_pages = await raw_get_scene_pagination_numbers(page)
    target_pages = [page_no for page_no in available_pages if page_no <= max_live_pages]
    rows = []
    reached_older_date = False

    for page_no in target_pages:
        await raw_switch_scene_page(page, page_no)
        live_rows = await raw_extract_main_live_rows(page, page_no)
        if max_live_rows > 0:
            live_rows = live_rows[:max_live_rows]
        emit_progress(
            progress,
            f"{shop_name}：读取第 {page_no} 页，共 {len(live_rows)} 场直播",
            kind="page",
            shop_name=shop_name,
            page=page_no,
        )

        for row_index, live_row in enumerate(live_rows):
            live_row["店铺名称"] = shop_name
            live_row["店铺网址"] = shop_url
            live_date = history_store.parse_live_date(live_row.get("直播时间", ""))

            if end_date and live_date and live_date > end_date:
                continue
            if start_date and live_date and live_date < start_date:
                reached_older_date = True
                break
            if (
                not force
                and history_store.is_historical(live_date)
                and history_store.is_session_complete(shop_name, live_row)
            ):
                emit_progress(
                    progress,
                    f"{shop_name} {live_row.get('直播时间', '')}：历史数据已存在，跳过",
                    kind="cached",
                    shop_name=shop_name,
                    live_date=live_date,
                )
                continue

            try:
                product_rows = await raw_capture_live_products(page, live_row, row_index)
            except Exception as exc:
                append_debug(
                    "raw capture live products failed: shop={} page={} row={} live={} error={}".format(
                        shop_name,
                        page_no,
                        row_index + 1,
                        live_row.get("直播时间", ""),
                        exc,
                    )
                )
                await raw_close_related_product_modal(page)
                emit_progress(
                    progress,
                    f"{shop_name} {live_row.get('直播时间', '')}：抓取失败，稍后可重试",
                    kind="error",
                    shop_name=shop_name,
                )
                continue

            history_store.save_session(shop_name, shop_url, live_row, product_rows)
            emit_progress(
                progress,
                f"{shop_name} {live_row.get('直播时间', '')}：新增 {len(product_rows)} 条成交商品",
                kind="captured",
                shop_name=shop_name,
                added=len(product_rows),
            )
            rows.extend(product_rows)

        if reached_older_date:
            break

    return rows


async def raw_capture_scene_products(cdp_url):
    ws_url = get_raw_cdp_page_ws(cdp_url)
    shop_name = get_arg_value("--shop-name", "")
    start_date = get_arg_value("--start-date", "")
    end_date = get_arg_value("--end-date", "")
    history_store.import_legacy_json_if_needed()
    async with RawCdpPage(ws_url) as page:
        await raw_capture_scene_products_from_page(
            page,
            shop_name=shop_name,
            start_date=start_date,
            end_date=end_date,
            force="--force" in sys.argv,
        )

    rows = history_store.query_records(start_date, end_date, [shop_name] if shop_name else None)
    save_scene_products(rows)
    print("scene product rows:", len(rows), flush=True)
    print("scene product json:", SCENE_PRODUCTS_JSON, flush=True)
    print("scene product csv:", SCENE_PRODUCTS_CSV, flush=True)
    print("scene product html:", SCENE_PRODUCTS_HTML, flush=True)
    print("scene product xlsx:", SCENE_PRODUCTS_XLSX, flush=True)
    return rows


async def raw_capture_scene_products_for_shops(
    cdp_url,
    shops,
    start_date="",
    end_date="",
    max_live_pages=None,
    force=False,
    progress=None,
    export_after=True,
):
    ws_url = get_raw_cdp_page_ws(cdp_url, allow_any_page=True)
    history_store.import_legacy_json_if_needed()
    selected_shop_names = [shop["name"] for shop in shops]

    async with RawCdpPage(ws_url) as page:
        for index, shop in enumerate(shops, 1):
            emit_progress(
                progress,
                f"店铺 {index}/{len(shops)}：{shop['name']}",
                kind="shop",
                shop_name=shop["name"],
            )
            try:
                await raw_navigate_to_url(page, shop["url"])
                await raw_capture_scene_products_from_page(
                    page,
                    shop_name=shop["name"],
                    shop_url=shop["url"],
                    start_date=start_date,
                    end_date=end_date,
                    max_live_pages=max_live_pages,
                    force=force,
                    progress=progress,
                )
            except Exception as exc:
                append_debug(f"shop capture failed: {shop['name']} {shop['url']} {exc}")
                emit_progress(
                    progress,
                    f"{shop['name']}：抓取失败 - {exc}",
                    kind="error",
                    shop_name=shop["name"],
                )

    all_rows = history_store.query_records(start_date, end_date, selected_shop_names)
    if export_after:
        save_scene_products(all_rows)
        print("scene product rows:", len(all_rows), flush=True)
        print("scene product json:", SCENE_PRODUCTS_JSON, flush=True)
        print("scene product csv:", SCENE_PRODUCTS_CSV, flush=True)
        print("scene product html:", SCENE_PRODUCTS_HTML, flush=True)
        print("scene product xlsx:", SCENE_PRODUCTS_XLSX, flush=True)
    return all_rows


def capture_scene_detail(page):
    open_scene_analysis(page)
    page_numbers = get_scene_pagination_numbers(page)
    rows = []

    for page_no in page_numbers:
        if len(page_numbers) > 1:
            switch_scene_page(page, page_no)
        page_rows = extract_current_scene_table(page, page_no)
        print(f"scene page {page_no}: {len(page_rows)} rows")
        rows.extend(page_rows)

    save_scene_detail(rows)
    print("scene detail rows:", len(rows))
    print("scene detail json:", SCENE_DETAIL_JSON)
    print("scene detail csv:", SCENE_DETAIL_CSV)
    return rows


def main():
    reset_profile_if_requested()

    if "--products-xlsx-only" in sys.argv:
        if not SCENE_PRODUCTS_JSON.exists():
            raise FileNotFoundError(
                f"missing input json: {SCENE_PRODUCTS_JSON.resolve()}. "
                "Run the scene product crawler first, then rerun --products-xlsx-only."
            )

        rows = json.loads(SCENE_PRODUCTS_JSON.read_text(encoding="utf-8"))
        if not isinstance(rows, list):
            raise RuntimeError(f"input json should be a list, got {type(rows).__name__}")
        output_path = save_scene_products_xlsx(rows)
        if output_path is None:
            raise RuntimeError("xlsx not generated. Please install openpyxl and Pillow.")
        print("scene product xlsx:", output_path.resolve(), flush=True)
        return

    if "--raw-cdp" in sys.argv and "--scene-products" in sys.argv:
        cdp_url = get_arg_value("--cdp-url", CDP_URL)
        print("raw cdp:", cdp_url, flush=True)
        shop_list_path = get_arg_value("--shop-list")
        single_url = get_arg_value("--url")
        start_date = get_arg_value("--start-date", "")
        end_date = get_arg_value("--end-date", "")
        max_live_pages = get_arg_int("--max-pages", 2)
        force = "--force" in sys.argv
        if shop_list_path:
            shops = parse_shop_list(shop_list_path)
            asyncio.get_event_loop().run_until_complete(
                raw_capture_scene_products_for_shops(
                    cdp_url,
                    shops,
                    start_date=start_date,
                    end_date=end_date,
                    max_live_pages=max_live_pages,
                    force=force,
                )
            )
        elif single_url:
            shops = [
                {
                    "name": get_arg_value("--shop-name", ""),
                    "url": single_url,
                }
            ]
            asyncio.get_event_loop().run_until_complete(
                raw_capture_scene_products_for_shops(
                    cdp_url,
                    shops,
                    start_date=start_date,
                    end_date=end_date,
                    max_live_pages=max_live_pages,
                    force=force,
                )
            )
        else:
            asyncio.get_event_loop().run_until_complete(raw_capture_scene_products(cdp_url))
        return

    with sync_playwright() as p:
        cdp_mode = "--cdp" in sys.argv
        scene_detail_mode = "--scene-detail" in sys.argv
        scene_products_mode = "--scene-products" in sys.argv
        cdp_browser = None

        if cdp_mode:
            cdp_url = get_arg_value("--cdp-url", CDP_URL)
            print("connect:", cdp_url)
            print("debug log:", DEBUG_LOG)
            cdp_browser = p.chromium.connect_over_cdp(cdp_url, timeout=60000)
            browser = cdp_browser.contexts[0]
        else:
            # 持久化浏览器资料，第一次手动扫码/登录，后面复用登录态
            executable_path = find_system_browser()
            launch_kwargs = {
                "user_data_dir": str(PROFILE),
                "headless": False,
                "no_viewport": True,
                "locale": "zh-CN",
                "timezone_id": "Asia/Shanghai",
                "args": [
                    "--start-maximized",
                    "--disable-blink-features=AutomationControlled",
                ],
                "ignore_default_args": ["--enable-automation"],
            }
            if executable_path:
                launch_kwargs["executable_path"] = executable_path

            print("browser:", executable_path or "Playwright Chromium")
            print("profile:", PROFILE)
            print("debug log:", DEBUG_LOG)

            # 持久化浏览器资料。优先使用系统 Chrome/Edge，登录页通常比 Playwright Chromium 稳定。
            browser = p.chromium.launch_persistent_context(**launch_kwargs)
            browser.add_init_script(
                """
                Object.defineProperty(navigator, 'webdriver', {
                    get: () => undefined
                });
                """
            )

        page = browser.pages[-1] if browser.pages else browser.new_page()

        def save_response(resp):
            content_type = resp.headers.get("content-type", "")
            if "application/json" not in content_type:
                return

            # 先宽松抓取，之后再筛选真正的数据接口
            keywords = ["daren", "author", "profile", "servicehall", "buyin"]
            if not any(k in resp.url.lower() for k in keywords):
                return

            try:
                data = resp.json()
            except Exception as exc:
                append_debug(f"json parse failed: {resp.url} {exc}")
                return

            filename = OUT / f"{int(time.time() * 1000)}.json"
            filename.write_text(
                json.dumps(
                    {
                        "url": resp.url,
                        "data": data,
                    },
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
            print("saved:", filename)

        page.on("response", save_response)
        page.on("requestfailed", lambda req: append_debug(f"request failed: {req.url}"))
        page.on("pageerror", lambda err: append_debug(f"page error: {err}"))
        page.on("console", lambda msg: append_debug(f"console {msg.type}: {msg.text}"))

        if not (cdp_mode and "buyin.jinritemai.com" in page.url):
            try:
                page.goto(URL, wait_until="domcontentloaded", timeout=60000)
            except PlaywrightTimeoutError:
                append_debug(f"goto timeout: {URL}")
                print("页面打开超时，但浏览器已打开，可以继续手动登录。")
        elif scene_detail_mode or scene_products_mode:
            print("use current page:", page.url)

        if scene_products_mode:
            capture_scene_products(page)
            text = page.locator("body").inner_text(timeout=10000)
            Path("outputs/baiying_visible_text.txt").write_text(text, encoding="utf-8")
            if not cdp_mode:
                browser.close()
            return

        if scene_detail_mode:
            capture_scene_detail(page)
            text = page.locator("body").inner_text(timeout=10000)
            Path("outputs/baiying_visible_text.txt").write_text(text, encoding="utf-8")
            if not cdp_mode:
                browser.close()
            return

        print("")
        print("请在打开的浏览器里完成登录。")
        print("如果点击登录后一直加载：关闭浏览器，然后运行：python data_gather.py --reset-profile")
        print("如果仍然卡住，再运行：python data_gather.py --edge --reset-profile")

        # 第一次运行时，在打开的浏览器里手动登录
        input("登录并打开到数据详情页后，按回车继续...")

        # 如果页面上有“数据详情”tab，可以尝试点击
        try:
            page.get_by_text("数据详情", exact=True).click(timeout=5000)
            page.wait_for_timeout(3000)
        except Exception:
            pass

        # 兜底：保存当前页面可见文本
        text = page.locator("body").inner_text()
        Path("outputs/baiying_visible_text.txt").write_text(text, encoding="utf-8")

        print("done")
        if not cdp_mode:
            browser.close()

if __name__ == "__main__":
    main()
